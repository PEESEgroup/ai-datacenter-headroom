"""Reproduce source-matched PJM load evidence and the ERCOT scale comparison."""
from pathlib import Path
from datetime import date
import argparse
import hashlib
import json
import re

import fitz
import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
PARENTS=("AEP","APS","ATSI","BGE","COMED","DAYTON","DOM","DQE","JCPL","METED","PECO","PEPCO","PL","PS")
COLORS={2024:(.925,0,.549),2025:(0,.627,.89)}
OBSERVED_COLOR=(0,.333,.506)
PDF_POINT_ALLOWANCE=1.0


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path,value):
    path.write_text(json.dumps(value,indent=2,allow_nan=False)+"\n")


def canonical_sheet(name):
    return "DOM" if name=="DOM" else name


def read_workbook(path):
    book=load_workbook(path,read_only=True,data_only=True)
    records=[]
    for sheet in book:
        rows=list(sheet.values)
        assert rows[0]==("Zone","Year","Metric","Load MW")
        for row,values in enumerate(rows[1:],2):
            zone,year,metric,value=values
            assert isinstance(year,int) and isinstance(value,(int,float))
            vintage=re.fullmatch(r"(20\d\d) Load Forecast July Peak",metric)
            assert vintage or metric=="Historical July Peak"
            records.append(dict(sheet=sheet.title,source_zone=zone,year=year,metric=metric,
                vintage=int(vintage.group(1)) if vintage else 0,MW=float(value),
                row=row,value_cell=f"D{row}",level="RTO" if sheet.title=="RTO" else
                "parent_zone" if sheet.title in PARENTS else "subzone"))
    book.close()
    result=pd.DataFrame(records)
    assert not result.duplicated(["sheet","year","metric"]).any()
    return result


def july_comparisons(raw):
    actual=raw[raw.vintage.eq(0)].rename(columns={"MW":"actual_MW","value_cell":"actual_cell"})
    forecast=raw[raw.vintage.gt(0)&raw.year.ge(raw.vintage)].rename(
        columns={"MW":"forecast_MW","value_cell":"forecast_cell"})
    pairs=forecast.merge(actual[["sheet","year","actual_MW","actual_cell"]],
                         on=["sheet","year"],how="inner",validate="many_to_one")
    pairs["actual_minus_forecast_MW"]=pairs.actual_MW-pairs.forecast_MW
    pairs["difference_pct_of_forecast"]=100*pairs.actual_minus_forecast_MW/pairs.forecast_MW.replace(0,np.nan)
    pairs["target_minus_vintage_years"]=pairs.year-pairs.vintage
    primary=pairs[pairs.level.eq("parent_zone")&pairs.year.eq(2025)&pairs.vintage.eq(2025)].copy()
    assert set(primary.sheet)=={"AEP","COMED","DOM","PS"}
    primary=primary.sort_values("sheet")
    primary["forecast_share_pct"]=100*primary.forecast_MW/primary.forecast_MW.sum()
    primary["actual_share_pct"]=100*primary.actual_MW/primary.actual_MW.sum()
    primary["share_change_pp"]=primary.actual_share_pct-primary.forecast_share_pct
    coverage=[]
    for zone in PARENTS:
        f=raw[raw.sheet.eq(zone)&raw.year.eq(2025)&raw.vintage.eq(2025)]
        a=raw[raw.sheet.eq(zone)&raw.year.eq(2025)&raw.vintage.eq(0)]
        coverage.append(dict(zone=zone,forecast_present=len(f)==1,actual_present=len(a)==1,
            matched=len(f)==len(a)==1,forecast_MW=float(f.MW.iloc[0]) if len(f) else None,
            actual_MW=float(a.MW.iloc[0]) if len(a) else None,
            treatment="matched" if len(f)==len(a)==1 else "not paired; missing is not zero"))
    return pairs,primary,pd.DataFrame(coverage)


def same_color(value,target):
    return value is not None and np.allclose(value,target,atol=.002,rtol=0)


def line_points(drawing):
    result=[]
    for item in drawing["items"]:
        assert item[0]=="l", "Expected vector polyline"
        result.extend([tuple(item[1]),tuple(item[2])])
    points=np.array(sorted(set(result)))
    # Shared segment endpoints can differ slightly in serialized coordinates.
    unique=[]
    for point in points:
        if unique and abs(point[0]-unique[-1][0])<.01:
            assert abs(point[1]-unique[-1][1])<.01
        else:
            unique.append(point)
    return np.array(unique)


def axis_calibration(page,drawings):
    grids=[d["rect"].y0 for d in drawings if d["type"]=="s"
           and same_color(d["color"],(.831,.831,.831))
           and d["rect"].width>400 and d["rect"].height<.01
           and 95<d["rect"].y0<270]
    ticks=[]
    for word in page.get_text("words"):
        x0,y0,x1,y1,label,*_=word
        if 55<x0<95 and 95<y0<267 and re.fullmatch(r"[1-9][\d,]*",label):
            nearest=min(grids,key=lambda y:abs(y-(y0+y1)/2))
            assert abs(nearest-(y0+y1)/2)<2
            ticks.append((nearest,float(label.replace(",",""))))
    assert len(ticks)>=2
    coordinates,values=np.array(ticks).T
    slope,intercept=np.polyfit(coordinates,values,1)
    fit_error=max(abs((values-intercept)/slope-coordinates))
    assert fit_error<PDF_POINT_ALLOWANCE
    return float(slope),float(intercept),ticks,float(fit_error)


def digitize_dc_months(path):
    document=fitz.open(path)
    records,calibrations,coverage=[],[],[]
    for page_index in range(9,len(document)):
        page=document[page_index]
        zone=page.get_text().splitlines()[0]
        if zone not in (*PARENTS,"RTO"):
            continue
        drawings=page.get_drawings()
        dots=[d for d in drawings if d["type"]=="f" and same_color(d["fill"],OBSERVED_COLOR)
              and 90<d["rect"].y0 and d["rect"].y1<270 and 7<d["rect"].width<10]
        paths={}
        for vintage,color in COLORS.items():
            found=[d for d in drawings if d["type"]=="s" and same_color(d["color"],color)
                   and d["rect"].width>150 and d["rect"].y0>90 and d["rect"].y1<270]
            assert len(found)<=1
            if found:
                paths[vintage]=line_points(found[0])
        coverage.append(dict(zone=zone,pdf_page=page_index+1,observed_months=len(dots),
                             forecast_vintages=",".join(map(str,paths))))
        if not dots:
            continue
        slope,intercept,ticks,fit_error=axis_calibration(page,drawings)
        assert 2025 in paths
        reference=paths[2025]
        dates=pd.date_range("2025-01-01",periods=len(reference),freq="MS")
        ordinals=np.array([d.date().toordinal() for d in dates])
        x_slope,x_intercept=np.polyfit(ordinals,reference[:,0],1)
        x_residual=max(abs((reference[:,0]-x_intercept)/x_slope-ordinals))
        assert x_residual<.05, "Forecast points must follow calendar-month dates"
        calibrations.append(dict(zone=zone,pdf_page=page_index+1,y_ticks=ticks,
            MW_per_point=slope,MW_intercept=intercept,max_tick_fit_error_points=fit_error,
            x_points_per_day=float(x_slope),x_intercept=float(x_intercept),
            max_date_fit_error_days=float(x_residual),coordinate_allowance_points=PDF_POINT_ALLOWANCE))
        for drawing in sorted(dots,key=lambda d:d["rect"].x0):
            rect=drawing["rect"]
            x,y=(rect.x0+rect.x1)/2,(rect.y0+rect.y1)/2
            ordinal=(x-x_intercept)/x_slope
            closest=min(pd.date_range("2024-01-01","2025-12-01",freq="MS"),
                        key=lambda d:abs(d.date().toordinal()-ordinal))
            assert abs(closest.date().toordinal()-ordinal)<.05
            observed=slope*y+intercept
            assert observed>0
            for vintage,points in paths.items():
                if closest.year<vintage:
                    continue
                index=int(np.argmin(abs(points[:,0]-x)))
                assert abs(points[index,0]-x)<.01
                prediction=slope*points[index,1]+intercept
                difference=observed-prediction
                allowance=2*abs(slope)*PDF_POINT_ALLOWANCE
                records.append(dict(zone=zone,month=closest.strftime("%Y-%m"),vintage=vintage,
                    observed_DC_peak_MW=observed,forecast_DC_peak_MW=prediction,
                    actual_minus_forecast_MW=difference,
                    difference_pct_of_forecast=100*difference/prediction,
                    difference_allowance_MW=allowance,
                    sign_with_allowance="above" if difference>allowance else
                    "below" if difference< -allowance else "unresolved_at_reading_allowance",
                    pdf_page=page_index+1,point_x=x,observed_point_y=y,
                    forecast_point_y=float(points[index,1]),digitized_from_vector_graph=True))
    document.close()
    frame=pd.DataFrame(records).sort_values(["zone","vintage","month"])
    assert set(frame.zone)=={"AEP","COMED","DOM","PS","RTO"}
    assert not frame.duplicated(["zone","vintage","month"]).any()
    summaries=[]
    common_window=frame[frame.month.between("2025-01","2025-08")]
    for (zone,vintage),group in common_window.groupby(["zone","vintage"]):
        counts=group.sign_with_allowance.value_counts()
        summaries.append(dict(zone=zone,vintage=int(vintage),months=len(group),
            first_month=group.month.min(),last_month=group.month.max(),
            mean_difference_MW=float(group.actual_minus_forecast_MW.mean()),
            MAE_MW=float(group.actual_minus_forecast_MW.abs().mean()),
            above=int(counts.get("above",0)),below=int(counts.get("below",0)),
            unresolved=int(counts.get("unresolved_at_reading_allowance",0)),
            difference_allowance_MW=float(group.difference_allowance_MW.max())))
    return frame,pd.DataFrame(summaries),calibrations,pd.DataFrame(coverage)


def extract_ercot_scale(page_pdf):
    document=fitz.open(page_pdf)
    assert len(document)==1
    table=next(t.extract() for t in document[0].find_tables().tables
               if t.extract()[0][:3]==["Summer","Total LL (MW)","LL Curtailment (MW)"])
    rows=[]
    for values in table[1:]:
        year=int(values[0].split()[0])
        clean=lambda value:float(value.replace(",","").replace(" ",""))
        rows.append(dict(year=year,total_large_load_MW=clean(values[1]),
                         provisional_curtailment_MW=clean(values[2])))
    document.close()
    return pd.DataFrame(rows)


def main(output):
    output.mkdir(parents=True,exist_ok=True)
    manifest=HERE/"INPUT_MANIFEST.json"
    if manifest.exists():
        for name,digest in json.loads(manifest.read_text())["files"].items():
            assert sha(HERE/name)==digest,"Input bytes differ: "+name
    raw=read_workbook(HERE/"sources/accuracy/2026-accuracy-report.xlsx")
    pairs,primary,coverage=july_comparisons(raw)
    months,summary,calibrations,month_coverage=digitize_dc_months(
        HERE/"sources/accuracy/2026-load-forecast-accuracy-report.pdf")
    ercot=extract_ercot_scale(HERE/"sources/accuracy/NERC_2025_ERCOT_p131.pdf")
    native=pd.read_csv(HERE/"inputs/native_inputs.csv")
    row=native[native.planning_region.eq("ERCOT")&native.year.eq(2035)].iloc[0]
    k=1+row.reference_margin_level_pct/100
    required=(k*(row.gross_peak_demand_MW-row.ordinary_response_MW)-row.anticipated_native_MW)/k
    ll=ercot[ercot.year.eq(2035)].iloc[0]
    np.testing.assert_allclose(ll.provisional_curtailment_MW,row.provisional_large_load_response_MW,atol=1e-7,rtol=0)
    scale=dict(total_large_load_GW_2035=float(ll.total_large_load_MW/1000),
        provisional_credit_GW_2035=float(ll.provisional_curtailment_MW/1000),
        required_effective_reduction_GW_2035=float(required/1000),
        required_pct_of_total_large_load=float(100*required/ll.total_large_load_MW),
        credit_pct_of_total_large_load=float(100*ll.provisional_curtailment_MW/ll.total_large_load_MW),
        required_pct_of_provisional_credit=float(100*required/ll.provisional_curtailment_MW))
    matched=dict(zones=primary.sheet.tolist(),forecast_total_MW=float(primary.forecast_MW.sum()),
        actual_total_MW=float(primary.actual_MW.sum()),
        actual_minus_forecast_MW=float(primary.actual_minus_forecast_MW.sum()),
        sum_absolute_zonal_difference_MW=float(primary.actual_minus_forecast_MW.abs().sum()),
        difference_pct_of_matched_forecast=float(100*primary.actual_minus_forecast_MW.sum()/primary.forecast_MW.sum()),
        covered_forecast_share_of_RTO_pct=float(100*primary.forecast_MW.sum()/raw[
            raw.sheet.eq("RTO")&raw.vintage.eq(2025)&raw.year.eq(2025)].MW.iloc[0]),
        coincident_system_peak_interpretation=False,DC_only_annual_interpretation=False)
    for name,frame in [("source_july_rows",raw),("all_available_july_pairs",pairs),
        ("july_2025_matched_parent_zones",primary),("july_coverage",coverage),
        ("DC_monthly_vector_pairs",months),("DC_monthly_summary",summary),
        ("DC_monthly_coverage",month_coverage),("ercot_large_load_source",ercot)]:
        frame.to_csv(output/(name+".csv"),index=False)
    write_json(output/"vector_calibration.json",calibrations)
    write_json(output/"ercot_large_load_scale.json",scale)
    write_json(output/"matched_july_summary.json",matched)
    source_urls={
        "2026-accuracy-report.xlsx":"https://www.pjm.com/-/media/DotCom/planning/res-adeq/load-forecast/2026-accuracy-report.xlsx",
        "2026-load-forecast-accuracy-report.pdf":"https://www.pjm.com/-/media/DotCom/planning/res-adeq/load-forecast/2026-load-forecast-accuracy-report.pdf",
        "NERC_2025_ERCOT_p131.pdf":"https://www.nerc.com/globalassets/our-work/assessments/nerc_ltra_2025.pdf"}
    write_json(output/"source_register.json",dict(retrieved="2026-10-23",sources=[
        dict(file="sources/accuracy/"+name,url=url,sha256=sha(HERE/"sources/accuracy"/name))
        for name,url in source_urls.items()],
        NERC_excerpt="Unmodified original page 131; original report hash in sources/accuracy/NERC_SOURCE.json"))
    write_json(output/"verification.json",dict(passed=True,workbook_rows=len(raw),
        July_primary_pairs=len(primary),DC_monthly_pairs=len(months),
        monthly_evidence="DC-only top charts, approximate vector recovery; one PDF point per ordinate",
        July_evidence="Exact workbook July large-load adjustments; matched parent-zone coverage only",
        missing_as_zero=False,subzones_added_to_parents=False,
        historical_error_extrapolated_to_2035=False,main_model_rerun=False))
    print(json.dumps(dict(July=matched,ERCOT=scale,monthly=summary.to_dict("records")),indent=2))


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output",type=Path,default=HERE/"empirical_results")
    main(parser.parse_args().output)
